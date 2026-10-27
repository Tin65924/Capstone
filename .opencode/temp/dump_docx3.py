import docx

d = docx.Document(r"C:\Users\astro\OneDrive\Desktop\Capstone\Laboratories\PEIS002_Lab12_Group404NotFound_MCST-LMS_Final_IT_Audit_Report.docx")

for i, t in enumerate(d.tables):
    txt = "\n".join(" | ".join(c.text.strip() for c in r.cells) for r in t.rows)
    for key in ("RSK-01", "CMP-01", "CRT-01", "REC-01", "EVD-09"):
        if key in txt:
            print(f"=== TABLE {i} (has {key}) ===")
            print(txt[:1800])
            print()
            break

# executive summary paragraph text anywhere
for para in d.paragraphs:
    t = para.text.strip()
    if t.startswith("The engagement") or "404NotFound performed" in t or "Overall conclusion" in t.lower() or t.startswith("The audit concluded"):
        print("PARA:", t[:500])
