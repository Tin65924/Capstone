from pptx import Presentation

p = Presentation(r"C:\Users\astro\OneDrive\Desktop\Capstone\Laboratories\MCST-LMS_Audit_Redesigned.pptx")
for i, s in enumerate(p.slides, 1):
    print(f"===== SLIDE {i} =====")
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                print(t)
        if sh.has_table:
            tbl = sh.table
            for r in tbl.rows:
                print(" | ".join(c.text.strip() for c in r.cells))
    if s.has_notes_slide:
        n = s.notes_slide.notes_text_frame.text.strip()
        print(f"--- notes --- {'(empty)' if not n else n[:300]}")
    print()
