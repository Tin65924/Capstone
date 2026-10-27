import docx

d = docx.Document(r"C:\Users\astro\OneDrive\Desktop\Capstone\Laboratories\PEIS002_Lab12_Group404NotFound_MCST-LMS_Final_IT_Audit_Report.docx")
lines = []
for para in d.paragraphs:
    t = para.text.strip()
    if t and para.style.name.startswith('Heading'):
        lines.append(f"## {para.style.name}: {t}")
print("\n".join(lines[:200]))
print("---TABLES---", len(d.tables))
for i, t in enumerate(d.tables[:40]):
    first = " | ".join(c.text.strip() for c in t.rows[0].cells)[:150]
    print(f"T{i}: rows={len(t.rows)} :: {first}")
