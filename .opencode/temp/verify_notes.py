import sys, io
from pptx import Presentation
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
p = Presentation(r"C:\Users\astro\OneDrive\Desktop\Capstone\Laboratories\MCST-LMS_Audit_Redesigned.pptx")
sl = list(p.slides)
for idx in (0, 10):
    s = sl[idx]
    print(f"===== SLIDE {idx+1} NOTES =====")
    for para in s.notes_slide.notes_text_frame.paragraphs:
        print(para.text[:120])
    print()
counts = []
for i, s in enumerate(sl, 1):
    lines = [q.text for q in s.notes_slide.notes_text_frame.paragraphs if q.text.strip()]
    qa = [l for l in lines if l and l[0].isdigit()]
    counts.append((i, len(qa)))
print("Q&A answers per slide:", counts)
